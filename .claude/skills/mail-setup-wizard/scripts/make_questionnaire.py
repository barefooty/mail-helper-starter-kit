# -*- coding: utf-8 -*-
"""產生《工作脈絡問卷.docx》（信箱小幫手起始包，第 0 題（怎麼答）選「問卷」時用）。

用法：
    python make_questionnaire.py --out 我的工作脈絡/工作脈絡問卷.docx
    python make_questionnaire.py --out 我的工作脈絡/工作脈絡問卷.docx --prefill 預填.json

預填 JSON 的鍵都可省略（格式見 references/產出規格.md 第 0b 項）：
    source, q3_lines, q5_rows, q6_skip, q6_urgent

版面規則（照 MEMORY：要她填的東西給 Word；必答紅字＋黃框、選填灰字；答案接受片段）：
    必答題：標題紅字、答案框淡黃底、黃框線
    選填題：標題灰字、答案框淡灰底
    全文微軟正黑體；標題用「Q1｜…」當錨點，read_questionnaire.py 靠它切題，不要改格式。
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "微軟正黑體"
RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
INK = RGBColor(0x2B, 0x3A, 0x45)
TEAL = RGBColor(0x2F, 0x6F, 0x73)
YELLOW_FILL = "FFF7CC"
YELLOW_LINE = "E0B000"
GRAY_FILL = "F2F2F2"
GRAY_LINE = "BFBFBF"
BLUE_FILL = "E9F2F2"
HEAD_FILL = "FBEFC8"      # 必答表格標題列
HEAD_FILL_OPT = "E7E7E7"  # 選填表格標題列
BOX = "☐"


# ---------- 基本工具 ----------
def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def para(doc, text="", size=11, bold=False, color=None, italic=False, space_after=4, align=None):
    p = doc.add_paragraph()
    if text:
        set_run_font(p.add_run(text), size, bold, color, italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    return p


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def borders(cell, color, size=8):
    tcpr = cell._element.get_or_add_tcPr()
    tcb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcb.append(el)
    tcpr.append(tcb)


def row_height(row, cm):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trpr.append(h)


def cell_text(cell, text, size=10, bold=False, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(text), size, bold, color, italic)


def style_all_cells(table, fill, line):
    for row in table.rows:
        for cell in row.cells:
            shade(cell, fill)
            borders(cell, line)
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, r.font.size.pt if r.font.size else 10)


# ---------- 問卷元件 ----------
def heading(doc, qno, title, required=True):
    """錨點格式固定為「Q1｜標題」，後面接 ［必答］／［選填］。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(f"Q{qno}｜{title}"), 13, True, RED if required else GRAY)
    set_run_font(p.add_run("　［必答］" if required else "　［選填，有填更準］"), 10, False, RED if required else GRAY)
    return p


def hint(doc, text):
    p = para(doc, text, 10, color=GRAY, space_after=4)
    p.paragraph_format.keep_with_next = True  # 提示別跟答案框分頁
    return p


def answer_box(doc, required=True, lines=3, placeholder=""):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    cell_text(c, placeholder, 10, color=GRAY, italic=True) if placeholder else cell_text(c, "")
    shade(c, YELLOW_FILL if required else GRAY_FILL)
    borders(c, YELLOW_LINE if required else GRAY_LINE, 12)
    row_height(t.rows[0], 0.75 * lines)
    para(doc, "", space_after=2)
    return t


def options(doc, items, note=None):
    """勾選題：每個選項一行「☐ 選項」。使用者把 ☐ 改成 ☑ 或在後面打 V。"""
    for it in items:
        p = para(doc, f"{BOX} {it}", 11, space_after=1)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.keep_with_next = True
    if note:
        hint(doc, note)


def grid(doc, headers, rows, required=True, prefill=None, prefill_note="側寫預填", widths=None):
    """表格題。prefill 是 list[dict]（鍵＝欄名），放在前幾列並在最後一欄註明來源。"""
    cols = len(headers)
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, 10, True, INK)
        shade(t.rows[0].cells[i], HEAD_FILL if required else HEAD_FILL_OPT)
    prefill = prefill or []
    for item in prefill:
        r = t.add_row()
        for i, h in enumerate(headers):
            v = str(item.get(h, "") or "")
            if i == cols - 1 and not v:
                v = prefill_note
            cell_text(r.cells[i], v, 10, color=INK if i < cols - 1 else GRAY, italic=(i == cols - 1))
        row_height(r, 0.8)
    for _ in range(max(rows - len(prefill), 2)):
        r = t.add_row()
        for c in r.cells:
            cell_text(c, "")
        row_height(r, 0.8)
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            borders(c, YELLOW_LINE if required else GRAY_LINE)
            if ri > 0:
                shade(c, YELLOW_FILL if required else GRAY_FILL)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    para(doc, "", space_after=2)
    return t


# ---------- 主體 ----------
def build(out: Path, prefill: dict):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)

    src = prefill.get("source")

    # 封面區
    para(doc, "信箱小幫手 起始包", 10, color=TEAL, space_after=0)
    para(doc, "工作脈絡問卷", 22, True, INK, space_after=2)
    para(doc, "填寫人：＿＿＿＿＿＿　　填寫日期：＿＿＿＿年＿＿月＿＿日", 11, space_after=6)
    if src:
        para(doc, f"預填來源：{src}。標「側寫預填」的內容是 Claude 從信件統計猜的，請直接改、刪、補。", 10, color=GRAY, space_after=6)

    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    c.text = ""
    lines = [
        ("怎麼填（看完這三行就好）", True),
        ("1. 紅字＋黃框的是必答（8 題），灰字是選填。片段、關鍵字、口語都可以，不用寫完整句，Claude 會幫你整理。", False),
        ("2. 勾選題把 ☐ 改成 ☑，或在選項後面打 V。表格不夠填就自己加列；第 4 題（專案階段）、第 5 題（對接對象）超過 10 列，改填同資料夾的兩張 Excel，這裡留空。", False),
        ("3. 填完存回「我的工作脈絡」資料夾（用 Cowork 或網頁版就直接上傳這個檔），回到 Claude 說「讀問卷」。沒填的題它會再問你，不會重問整份。", False),
    ]
    for i, (txt, b) in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(txt), 10, b, TEAL if b else INK)
    shade(c, BLUE_FILL)
    borders(c, "BFD6D6", 8)
    para(doc, "", space_after=4)

    # Q1
    heading(doc, 1, "你是誰、工作主要在做什麼")
    hint(doc, "職稱、單位、工作一句話。例：「計畫專員，承辦兩個跨校計畫」。平常寄信給你的，比較多是學生、校內同事、校外單位，還是廠商？")
    answer_box(doc, lines=3)

    # Q2
    heading(doc, 2, "要整理哪個信箱、誰在用")
    options(doc, ["我自己的工作信箱（一個人用）", "多人共用的信箱（幾個人、各負責什麼、報告要寄給誰，請寫在下面）", "兩個都有（先做一個，另一個之後複製設定）"])
    options(doc, ["這個信箱是 Google 帳號（Gmail 或學校的 Google Workspace）", "不是，或不確定"],
            note="目前兩個版本都只支援 Google 信箱；不確定的話先填，Claude 會跟你確認。")
    answer_box(doc, lines=2, placeholder="信箱地址：　　　　　　　　　　　　　　／共用的話誰負責什麼：")

    # Q3
    heading(doc, 3, "手上有哪幾條業務線、各自的週期")
    hint(doc, "業務線＝你工作的大分類，之後會變成報告的分區與標籤名稱。超過 6 條的話，先寫信最多的 3 條，其他可以合併成「其他」。參考：學生服務型（選課／成績／學籍／證明／獎學金）、計畫窗口型（一個計畫一條線）、採購總務型（請購／驗收／盤點／修繕）、人事主計型（差假／報帳／薪資／法規）、服務台型（報修／借用／帳號）。")
    q3_pre = [{"業務線名稱": n, "週期（學期制／年度制／專案制／常態）": "", "這條線信最多？（打 V）": "", "備註": ""} for n in prefill.get("q3_lines", [])]
    grid(doc, ["業務線名稱", "週期（學期制／年度制／專案制／常態）", "這條線信最多？（打 V）", "備註"], 6, True, q3_pre, widths=[5.0, 5.5, 3.5, 3.0])

    # Q4
    heading(doc, 4, "每條業務線會經過哪些階段、該收到什麼", required=False)
    hint(doc, "以一條線為例：一個學期（或一年）會經過哪些階段？大約什麼時候、會收到什麼信、來自誰？「該收到什麼」＋「大約時間」就是小幫手的「該來沒來」提醒。有整理好的表就改填《專案階段表.xlsx》，這裡留空。")
    grid(doc, ["業務線", "階段順序", "階段名稱", "大約時間", "會收到什麼信", "來自誰", "最容易出事的點"], 8, False, widths=[2.6, 1.4, 2.4, 2.0, 3.2, 2.2, 3.2])
    hint(doc, "哪個階段最容易出事、最需要盯？（那個階段的信會排前面）")
    answer_box(doc, required=False, lines=2)

    # Q5
    heading(doc, 5, "誰會寄信來、對接對象與他們的習慣")
    hint(doc, "一列一個單位或人。至少填「單位」與「email 或網域」，小幫手靠它認出誰寄的信。「習慣／注意」寫你希望被提醒的一句話，例：來信要固定 CC 某人、常漏附件、常用私人信箱寄。超過 10 列改填《對接對象清單.xlsx》。")
    q5_headers = ["單位", "聯絡人", "email 或網域", "所屬業務線", "習慣／注意（希望被提醒的一句）", "備註"]
    q5_pre = []
    for r in prefill.get("q5_rows", []):
        q5_pre.append({"單位": r.get("單位", ""), "聯絡人": r.get("聯絡人", ""), "email 或網域": r.get("email或網域", r.get("email 或網域", "")),
                       "所屬業務線": r.get("業務線", ""), "習慣／注意（希望被提醒的一句）": r.get("習慣／注意", ""), "備註": ""})
    grid(doc, q5_headers, 8, True, q5_pre, widths=[3.2, 1.8, 3.4, 2.6, 3.6, 2.4])
    hint(doc, "平常寄信給你的類型（可複選）：")
    options(doc, ["學生", "校內同仁", "校外機構窗口", "廠商", "系統通知（表單、平台、討論區）", "主管"])
    hint(doc, "有沒有哪些寄件人的信，一來就要優先看？")
    answer_box(doc, lines=2)

    # Q6
    heading(doc, 6, "哪些信最想先看、哪些一句帶過、什麼算急件")
    hint(doc, "6-1 收到信你第一個想知道什麼？哪些信要排最前面？（可複選＋補充）")
    options(doc, ["還沒回的", "有期限的", "特定對象寄的（寫誰）", "有附件要處理的"])
    answer_box(doc, lines=2)
    hint(doc, "6-2 哪些信一句帶過就好？（可複選＋補充）")
    skip_items = ["電子報、廣告", "系統通知、表單自動回覆", "已回覆的串"]
    for s in prefill.get("q6_skip", []):
        skip_items.append(f"{s}（側寫預填）")
    options(doc, skip_items)
    answer_box(doc, lines=2)
    hint(doc, "6-3 什麼樣的信算急？（可複選＋補充）")
    urgent_items = ["主旨或內文有期限、截止日", "主管寄的", "特定窗口寄的（寫誰）", "有「緊急」「催」「退件」字樣"]
    for s in prefill.get("q6_urgent", []):
        urgent_items.append(f"{s}（側寫預填）")
    options(doc, urgent_items)
    answer_box(doc, lines=2)

    # Q7
    heading(doc, 7, "報告多久收到一次、幾點、怎麼送")
    options(doc, ["每天一次（早上 8:30）", "每天兩次（8:00 與 14:00）", "每天三次（8、12、17）", "其他時間（寫在下面）"])
    options(doc, ["週末也要", "週五附本週摘要"])
    hint(doc, "報告怎麼送？")
    options(doc, ["寄一封摘要信給我自己", "不用寄，我到 Claude 的排程紀錄看"])
    hint(doc, "收件人（只有我／團隊每人一份，寫信箱）與其他想法：")
    answer_box(doc, lines=2)
    hint(doc, "［選填］報告信的顏色風格：A 沉穩深青（預設）／B 溫暖棕／C 知性靛藍／D 極簡黑白")
    answer_box(doc, required=False, lines=1)

    # Q8
    heading(doc, 8, "小幫手可以動你的信箱到什麼程度")
    options(doc, ["完全唯讀：只讀信、只給報告，信箱完全不變",
                  "貼小幫手專用標籤（建議）：只用「小幫手/」開頭的全新標籤，不碰你手動的標籤、不標已讀；好處是在 Gmail 就能用標籤篩信"],
            note="兩個選項都不會替你回信、寄信、建草稿、刪信。選「貼標籤」時 Gmail 會自動多一個上層標籤「小幫手」，屬正常現象。")
    answer_box(doc, lines=1, placeholder="想補充的（例如：不要碰某個標籤、想用的待辦標籤名）：")
    hint(doc, "8-2［選填］有沒有哪類信敏感到連摘要都不要出現在報告？（例：人事、主計、學生個資）這類信在報告只會有一行「王O明（學校）｜主旨｜時間」，不寫摘要。沒有就留空。"
              "（不管有沒有填，報告裡的身分證字號、電話、學號、長數字、帳號密碼都會先換成「○」。）")
    answer_box(doc, required=False, lines=1)

    # Q9
    heading(doc, 9, "想用哪個版本")
    options(doc, ["簡易版（建議）：在 Claude 網頁版連上 Gmail、建一個排程，約 10 到 15 分鐘，只要瀏覽器",
                  "進階版：自己架程式在雲端跑，約 1 到 2 小時，要 GitHub 與 Python；適合要程式層級把關、自訂報告版型的人",
                  "不確定，請 Claude 建議"])

    # Q10
    heading(doc, 10, "其他想讓小幫手知道的", required=False)
    hint(doc, "例：固定週期一定會收到的信（每月底報帳通知、每學期各校名單）、不希望出現在報告的東西、你最怕漏掉的信。")
    answer_box(doc, required=False, lines=3)

    para(doc, "填完存回「我的工作脈絡」資料夾，回到 Claude 說「讀問卷」。這份問卷只留在你的電腦，不要上傳到公開的地方。", 9, color=GRAY, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description="產生工作脈絡問卷.docx")
    ap.add_argument("--out", required=True, help="輸出路徑（.docx）")
    ap.add_argument("--prefill", help="預填 JSON（UTF-8）")
    a = ap.parse_args()
    prefill = {}
    if a.prefill:
        prefill = json.loads(Path(a.prefill).read_text(encoding="utf-8"))
    out = build(Path(a.out), prefill)
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"已產生：{out}（預填：{'有' if prefill else '無'}）")


if __name__ == "__main__":
    main()
